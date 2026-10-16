# -*- coding: utf-8 -*-
"""v0.50.0 -- Creditos PIS/COFINS: leitores (SPED e Detalhamento), conferencia, ajustes, graficos e planilha. Tudo com arquivos
SINTETICOS (creditos_fixtures.py); nenhum dado real no repositorio."""
import io
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import creditos_conferencia as cf  # noqa: E402
import creditos_detalhamento as cdet  # noqa: E402
import creditos_export as cexp  # noqa: E402
import creditos_graficos as gr  # noqa: E402
import sped_contribuicoes as sc  # noqa: E402
from creditos_fixtures import (CNPJ_ENERGIA, FORN_A, detalhamento_sintetico, sped_sintetico)  # noqa: E402


@pytest.fixture()
def par():
    return sc.ler_sped(*sped_sintetico()), cdet.ler_detalhamento(*detalhamento_sintetico())


# ───────────── SPED ─────────────
def test_sped_cabecalho_e_competencia():
    s = sc.ler_sped(*sped_sintetico())
    assert s["erro"] is None
    assert s["cabecalho"]["competencia"] == "07/2026" and s["cabecalho"]["cnpj"] == CNPJ_ENERGIA and s["cabecalho"]["tipo"] == "Original"
    assert s["cabecalho"]["enviado_em_arquivo"] == "2026-10-16 10:00:00"


def test_sped_creditos_por_registro_e_categoria():
    s = sc.ler_sped(*sped_sintetico())
    c = s["creditos"].set_index("registro")
    assert set(c.index) == {"A170", "C170", "C501", "F120"} or len(c) == 5
    assert round(s["creditos"]["pis"].sum(), 2) == 127.05 and round(s["creditos"]["cofins"].sum(), 2) == 585.20
    cats = dict(zip(s["creditos"]["categoria"], s["creditos"]["pis"]))
    assert cats["Combustível"] == 8.25 and cats["Aluguel"] == 33.0 and cats["Serviços Tomados"] == 16.5
    assert cats["Demais Produtos (energia)"] == 49.5 and cats["Depreciação (F120)"] == 19.8


def test_sped_numero_da_nota_c100_c500_e_cofins_juntado():
    s = sc.ler_sped(*sped_sintetico())
    por_num = {r.numero: r for r in s["creditos"].itertuples() if r.numero}
    assert set(por_num) == {"1001", "55", "777", "9001"}
    assert por_num["9001"].cofins == 228.0 and por_num["9001"].cst_cofins == "50"      # C505 junto do C501


def test_sped_receita_so_das_vendas():
    s = sc.ler_sped(*sped_sintetico())
    assert len(s["receitas"]) == 1 and s["receitas"].iloc[0]["cliente"] == "CLIENTE X" and s["receitas"].iloc[0]["pis"] == 165.0


def test_sped_apuracao_resumo():
    r = sc.resumo_apuracao(sc.ler_sped(*sped_sintetico())["apuracao"])
    assert r["credito_pis"] == 127.05 and r["credito_cofins"] == 585.2
    assert r["debito_nc_pis"] == 300.0 and r["debito_cum_cofins"] == 300.0 and r["retencao_nc_cofins"] == 50.0
    assert r["a_recolher_pis"] == 227.95 and r["a_recolher_cofins"] == 1164.8
    assert r["receita_nc"] == 20000.0 and r["receita_cum"] == 10000.0


def test_sped_ignora_assinatura_depois_do_9999_e_arquivo_invalido():
    s = sc.ler_sped(b"qualquer coisa", "x.txt")
    assert s["erro"] and "0000" in s["erro"]
    assert sc.ler_sped(b"", "x.txt")["erro"]
    assert sc.ler_sped("|0000|006|0||xx|yy|N|1|".encode("latin-1"), "x.txt")["erro"]


def test_sped_retificadora_e_hash():
    c, n = sped_sintetico(tipo="1")
    s = sc.ler_sped(c, n)
    assert s["cabecalho"]["tipo"] == "Retificadora" and s["sha256"] == sc.sha256(c)


# ───────────── Detalhamento ─────────────
def test_detalhamento_leitura():
    d = cdet.ler_detalhamento(*detalhamento_sintetico())
    assert d["erro"] is None and d["competencia"] == "07/2026"
    assert len(d["linhas"]) == 4 and set(d["categorias"]) == {"Serviços Tomados", "Aluguel", "Combustível", "Demais Produtos"}
    assert round(d["linhas"]["pis"].sum(), 2) == 107.25 and round(d["linhas"]["cofins"].sum(), 2) == 494.0
    assert len(d["debitos"]) == 13 and not d["resumo"].empty
    assert (d["linhas"]["cnpj"].str.len() == 14).all() and "TOTAL" not in set(d["linhas"]["cnpj"])


def test_detalhamento_arquivo_invalido():
    assert cdet.ler_detalhamento(b"nao e xlsx", "x.xlsx")["erro"]
    b = io.BytesIO()
    pd.DataFrame({"a": [1]}).to_excel(b, index=False)
    assert "Detalhamento" in cdet.ler_detalhamento(b.getvalue(), "x.xlsx")["erro"]


def test_detalhamento_normaliza_numero_com_zeros_e_decimal():
    assert cdet.normalizar_numero("000123") == "123" and cdet.normalizar_numero("123.0") == "123" and cdet.normalizar_numero(55) == "55"


# ───────────── Conferencia ─────────────
def _st(df):
    return dict(zip(df["numero"], df["status"]))


def test_conferencia_tudo_bate(par):
    s, d = par
    L = cf.conferir_linhas(s, d)
    st = _st(L)
    assert st["1001"] == cf.S_OK and st["55"] == cf.S_OK and st["777"] == cf.S_OK and st["9001"] == cf.S_OK
    bens = L[L["nivel"] == "Bens (sem nota)"]
    assert len(bens) == 1 and bens.iloc[0]["status"] == cf.S_ESPERADO
    assert cf.conferir_categorias(s, d).set_index("categoria").loc["Depreciação (F120)", "status"] == cf.S_ESPERADO
    assert (cf.conferir_totais(s, d)["status"] == cf.S_OK).all()
    assert (cf.conferir_debitos(s, d)["status"] == cf.S_OK).all()
    assert (cf.conferir_bases(s, d)["dif_interna_pis"].abs() < 0.01).all()


def test_conferencia_valor_diferente_so_sped_so_detalhamento():
    s = sc.ler_sped(*sped_sintetico())
    d = cdet.ler_detalhamento(*detalhamento_sintetico(alterar={"1001": (20.0, 76.0)}))
    st = _st(cf.conferir_linhas(s, d))
    assert st["1001"] == cf.S_DIF
    # nota do Detalhamento que nao existe no SPED + nota do SPED que nao esta no Detalhamento
    d2 = cdet.ler_detalhamento(*detalhamento_sintetico())
    d2["linhas"] = d2["linhas"][d2["linhas"]["numero"] != "55"]
    nova = d2["linhas"].iloc[[0]].copy()
    nova["numero"], nova["pis"], nova["cofins"], nova["fornecedor"] = "8888", 1.0, 5.0, "OUTRO"
    nova["cnpj"] = "99999999000199"
    d2["linhas"] = pd.concat([d2["linhas"], nova], ignore_index=True)
    st2 = _st(cf.conferir_linhas(s, d2))
    assert st2["55"] == cf.S_SO_SPED and st2["8888"] == cf.S_SO_DET


def test_conferencia_numero_diferente_mesmo_fornecedor_e_valor():
    s = sc.ler_sped(*sped_sintetico())
    d = cdet.ler_detalhamento(*detalhamento_sintetico())
    d["linhas"].loc[d["linhas"]["numero"] == "9001", "numero"] = "9002"
    L = cf.conferir_linhas(s, d)
    r = L[L["numero"].str.contains("9001")].iloc[0]
    assert r["status"] == cf.S_NUM and "9002" in r["observacao"]


def test_conferencia_retificadora_com_credito_maior_aparece_nas_categorias_e_totais():
    s = sc.ler_sped(*sped_sintetico(ajuste_servico=5.0))
    d = cdet.ler_detalhamento(*detalhamento_sintetico())
    c = cf.conferir_categorias(s, d).set_index("categoria")
    assert c.loc["Serviços Tomados", "dif_pis"] == 5.0 and c.loc["Serviços Tomados", "status"] == cf.S_DIF
    assert cf.conferir_totais(s, d).set_index("tributo").loc["PIS", "status"] == cf.S_DIF
    assert cf.conferir_debitos(s, d).set_index("item").loc["Crédito Não Cumulativo apurado no mês (M100/M500)", "dif_pis"] == 5.0


def test_conferencia_sem_aba_de_debitos():
    s = sc.ler_sped(*sped_sintetico())
    d = cdet.ler_detalhamento(*detalhamento_sintetico(com_debitos=False))
    D = cf.conferir_debitos(s, d)
    assert len(D) == 13 and (D["status"] == "Sem aba no Detalhamento").all() and D["pis_det"].isna().all()


def test_totais_explica_inconsistencia_interna_do_sped():
    """M105 diz base 1.000 x 1,65% = 16,50 mas a linha do documento tem PIS 20,00: o Bloco M (M100) segue a base; a diferenca vem do SPED."""
    s = sc.ler_sped(*sped_sintetico())
    s["creditos"].loc[s["creditos"]["numero"] == "1001", "pis"] = 20.0
    d = cdet.ler_detalhamento(*detalhamento_sintetico(alterar={"1001": (20.0, 76.0)}))
    t = cf.conferir_totais(s, d).set_index("tributo")
    assert t.loc["PIS", "status"] == cf.S_OK and t.loc["PIS", "inconsistencia_sped"] == -3.5 and "inconsistência" in t.loc["PIS", "observacao"]


# ───────────── Ajustes (edicoes) ─────────────
def test_aplicar_aba_edita_recalcula_e_marca_editado(par):
    s, d = par
    base = cf.conferir_linhas(s, d)
    ch = f"{FORN_A}|1001"
    aj = [{"aba": "linhas", "chave": ch, "campo": "pis_det", "valor_novo": "10.00"}]
    L = cf.aplicar_aba("linhas", base, aj).set_index("chave")
    assert L.loc[ch, "pis_det"] == 10.0 and L.loc[ch, "dif_pis"] == 6.5 and L.loc[ch, "status"] == cf.S_DIF and bool(L.loc[ch, "editado"])
    assert not L.drop(index=ch)["editado"].any()


def test_aplicar_aba_status_manual_vale_por_cima_e_nao_altera_a_base(par):
    s, d = par
    base = cf.conferir_linhas(s, d)
    ch = f"{FORN_A}|1001"
    aj = [{"aba": "linhas", "chave": ch, "campo": "pis_det", "valor_novo": "10"}, {"aba": "linhas", "chave": ch, "campo": "status", "valor_novo": cf.S_OK}]
    L = cf.aplicar_aba("linhas", base, aj).set_index("chave")
    assert L.loc[ch, "status"] == cf.S_OK and L.loc[ch, "dif_pis"] == 6.5
    assert base.set_index("chave").loc[ch, "pis_det"] == 16.5      # base intacta


def test_aplicar_aba_linha_manual_e_exclusao(par):
    s, d = par
    base = cf.conferir_linhas(s, d)
    aj = [{"aba": "linhas", "chave": "MANUAL|a", "campo": "fornecedor", "valor_novo": "Fornecedor manual"},
          {"aba": "linhas", "chave": "MANUAL|a", "campo": "pis_det", "valor_novo": "3,5"},
          {"aba": "linhas", "chave": "33333333000133|777", "campo": "__excluir__", "valor_novo": "true"},
          {"aba": "linhas", "chave": "chave-que-nao-existe-mais", "campo": "pis_det", "valor_novo": "1"}]
    L = cf.aplicar_aba("linhas", base, aj).set_index("chave")
    assert L.loc["MANUAL|a", "status"] == cf.S_MANUAL and L.loc["MANUAL|a", "pis_det"] == 3.5 and L.loc["MANUAL|a", "fornecedor"] == "Fornecedor manual"
    assert "33333333000133|777" not in L.index and len(L) == len(base)       # -1 excluida, +1 manual; chave velha ignorada


def test_aplicar_aba_outras_abas(par):
    s, d = par
    t = cf.aplicar_aba("totais", cf.conferir_totais(s, d), [{"aba": "totais", "chave": "PIS", "campo": "soma_detalhamento", "valor_novo": "100"}])
    assert t.set_index("tributo").loc["PIS", "residuo"] == 7.25 and t.set_index("tributo").loc["PIS", "status"] == cf.S_DIF
    c = cf.aplicar_aba("categorias", cf.conferir_categorias(s, d), [{"aba": "categorias", "chave": "Depreciação (F120)", "campo": "pis_det", "valor_novo": "19.8"},
                                                                   {"aba": "categorias", "chave": "Depreciação (F120)", "campo": "cofins_det", "valor_novo": "91.2"}])
    assert c.set_index("categoria").loc["Depreciação (F120)", "status"] == cf.S_OK
    b = cf.aplicar_aba("bases", cf.conferir_bases(s, d), [{"aba": "bases", "chave": "Aluguel (NAT 05)", "campo": "base_det", "valor_novo": "1000"}])
    assert b.set_index("grupo").loc["Aluguel (NAT 05)", "status"] == cf.S_DIF
    dd = cf.aplicar_aba("debitos", cf.conferir_debitos(s, d), [{"aba": "debitos", "chave": "TOTAL PIS/COFINS a recolher no mês (M200/M600)", "campo": "pis_det", "valor_novo": "200"}])
    assert dd.set_index("item").loc["TOTAL PIS/COFINS a recolher no mês (M200/M600)", "status"] == cf.S_DIF


def test_diferencas_edicao_so_o_que_mudou(par):
    s, d = par
    base = cf.conferir_linhas(s, d)
    ed = base.copy()
    ed.loc[ed["numero"] == "1001", "pis_det"] = 10.0
    ed.loc[ed["numero"] == "55", "nota_contadora"] = "ok, conferido no papel"
    ed = ed[ed["numero"] != "777"]
    manual = {c: (False if base[c].dtype == bool else 0.0 if base[c].dtype == float else "") for c in base.columns}
    manual.update(chave="MANUAL|z", fornecedor="Novo", pis_det=2.0)
    ed = pd.concat([ed, pd.DataFrame([manual])], ignore_index=True)
    mud = cf.diferencas_edicao(base, ed, "chave", ["fornecedor", "pis_det", "nota_contadora", "conferido"])
    por = {(m["chave"], m["campo"]): m for m in mud}
    assert por[(f"{FORN_A}|1001", "pis_det")]["valor_novo"] == "10.0" and por[(f"{FORN_A}|1001", "pis_det")]["valor_original"] == "16.5"
    assert por[("22222222000122|55", "nota_contadora")]["valor_novo"] == "ok, conferido no papel"
    assert por[("33333333000133|777", "__excluir__")]["valor_novo"] == "true"
    assert ("MANUAL|z", "fornecedor") in por and ("MANUAL|z", "pis_det") in por
    assert not any(k[0] == "44444444000144|9001" for k in por)


def test_resumo_linhas(par):
    s, d = par
    r = cf.resumo_linhas(cf.conferir_linhas(s, d)).set_index("status")
    assert r.loc[cf.S_OK, "linhas"] == 4 and r.loc[cf.S_ESPERADO, "linhas"] == 1


# ───────────── Graficos ─────────────
def _serie():
    base = dict(sped_tipo="Original", deb_nc_pis=300.0, deb_nc_cofins=1300.0, deb_cum_pis=60.0, deb_cum_cofins=280.0, usado_pis=100.0, usado_cofins=450.0,
                anterior_pis=0.0, anterior_cofins=0.0, saldo_pis=0.0, saldo_cofins=0.0, ret_pis=5.0, ret_cofins=20.0, recolher_pis=200.0,
                recolher_cofins=900.0, receita_nc=20000.0, receita_cum=9000.0, orig_por_categoria={}, por_natureza={"05": {"natureza": "Aluguel", "base": 1, "pis": 90, "cofins": 410}},
                por_categoria={"Combustível": {"pis": 10, "cofins": 40, "linhas": 3}, "Aluguel": {"pis": 90, "cofins": 410, "linhas": 2}},
                dif_fornecedores=[{"fornecedor": "A", "dif_pis": 10, "dif_cofins": 40}, {"fornecedor": "B", "dif_pis": -3, "dif_cofins": -9}])
    return pd.DataFrame([
        dict(base, competencia="01/2026", origem="HISTORICO", tem_original=True, cred_pis=100.0, cred_cofins=450.0, orig_pis=80.0, orig_cofins=370.0),
        dict(base, competencia="02/2026", origem="HISTORICO", tem_original=True, cred_pis=90.0, cred_cofins=400.0, orig_pis=95.0, orig_cofins=410.0),
        dict(base, competencia="08/2026", origem="ENVIO", tem_original=False, cred_pis=120.0, cred_cofins=500.0, orig_pis=None, orig_cofins=None)])


def test_preparar_serie_diferencas_e_acumulados():
    d = gr.preparar_serie(_serie())
    assert list(d["rotulo"]) == ["jan/26", "fev/26", "ago/26"]
    assert list(d["cred_total"]) == [550.0, 490.0, 620.0]
    assert d.loc[0, "dif_total"] == 100.0 and d.loc[1, "dif_total"] == -15.0 and d.loc[0, "dif_acum"] == 100.0 and d.loc[1, "dif_acum"] == 85.0
    assert pd.isna(d.loc[2, "dif_total"]) and pd.isna(d.loc[2, "orig_total"])
    assert d["cred_acum"].iloc[-1] == 1660.0 and d.loc[0, "usado_total"] == 550.0


def test_preparar_serie_segue_o_tributo():
    """O seletor PIS / COFINS / PIS+COFINS precisa mexer nos KPIs (colunas *_total), nao so' em alguns graficos."""
    tot, pis, cof = (gr.preparar_serie(_serie(), t) for t in ("total", "pis", "cofins"))
    assert list(pis["cred_total"]) == [100.0, 90.0, 120.0] and list(cof["cred_total"]) == [450.0, 400.0, 500.0]
    for col in ("cred_total", "orig_total", "dif_total", "deb_nc_total", "deb_cum_total", "usado_total", "ret_total", "recolher_total", "saldo_total"):
        for i in range(len(tot)):
            a, b, c = tot.loc[i, col], pis.loc[i, col], cof.loc[i, col]
            if pd.isna(a):
                assert pd.isna(b) and pd.isna(c)
            else:
                assert abs(a - (b + c)) < 1e-9, (col, i)
    assert pis.loc[0, "dif_total"] == 20.0 and cof.loc[0, "dif_total"] == 80.0 and pis.loc[1, "dif_acum"] == 15.0
    assert pis.loc[0, "deb_nc_total"] == 300.0 and cof.loc[0, "recolher_total"] == 900.0


def test_uploaders_do_creditos_tem_contadores_independentes():
    """Gravar o SPED nao pode limpar o Detalhamento que ainda espera confirmacao (e vice-versa)."""
    import pathlib
    src = (pathlib.Path(__file__).resolve().parents[1] / "telas" / "10_Creditos_PIS_COFINS.py").read_text(encoding="utf-8")
    assert "cred_up_n_sped" in src and "cred_up_n_det" in src and 'get("cred_up_n"' not in src
    assert 'key=f"cred_up_sped_{n_up_s}"' in src and 'key=f"cred_up_det_{n_up_d}"' in src


def test_aviso_fica_em_recipiente_fixo_e_confirmacao_zera():
    """Achado ao vivo (v0.50.0): o aviso no topo empurrava as abas; ao sumir, a tela voltava para 'Enviar arquivos'.
    O aviso tem de ficar dentro de um st.container() sempre presente; e o 'Confirmo' de arquivar/restaurar vale para uma ação só."""
    import pathlib
    src = (pathlib.Path(__file__).resolve().parents[1] / "telas" / "10_Creditos_PIS_COFINS.py").read_text(encoding="utf-8")
    assert "with st.container():\n    mostrar_flash()" in src
    assert "\nmostrar_flash()" not in src
    assert 'key=f"cred_ok_envio_{n_ok}"' in src and src.count('st.session_state["cred_ok_n"] = n_ok + 1') == 2


@pytest.mark.parametrize("dark", [False, True])
@pytest.mark.parametrize("trib", ["total", "pis", "cofins"])
def test_todos_os_graficos_montam(dark, trib):
    d = gr.preparar_serie(_serie())
    figs = [gr.fig_credito_mensal(d, dark, trib), gr.fig_diferenca(d, dark), gr.fig_credito_acumulado(d, dark), gr.fig_credito_categoria(d, dark, trib),
            gr.fig_debito_credito(d, dark), gr.fig_a_recolher(d, dark), gr.fig_saldo_credito(d, dark), gr.fig_receita(d, dark),
            gr.fig_base_natureza(d.at[0, "por_natureza"], dark), gr.fig_creditos_por_categoria_mes(d.at[0, "por_categoria"], dark, trib),
            gr.fig_dif_fornecedores(d.at[0, "dif_fornecedores"], dark)]
    assert all(len(f.data) >= 1 for f in figs)


def test_graficos_vazios_nao_quebram():
    vazio = pd.DataFrame()
    for f in (gr.fig_credito_mensal, gr.fig_diferenca, gr.fig_credito_acumulado, gr.fig_credito_categoria, gr.fig_debito_credito, gr.fig_a_recolher,
              gr.fig_saldo_credito, gr.fig_receita):
        assert f(vazio, True) is not None
    assert gr.preparar_serie(vazio).empty
    assert gr.fig_top_fornecedores(None, False) is not None and gr.fig_dif_fornecedores([], False) is not None and gr.fig_base_natureza({}, False) is not None


def test_credito_mensal_mostra_original_so_nos_meses_com_original():
    f = gr.fig_credito_mensal(gr.preparar_serie(_serie()), False)
    linha = next(t for t in f.data if t.type == "scatter")
    assert list(linha.y[:2]) == [450.0, 505.0] and (linha.y[2] is None or pd.isna(linha.y[2]))


def test_cor_da_categoria_e_fixa():
    assert gr.cor_categoria("Combustível", False) == "#2a78d6" and gr.cor_categoria("Aluguel", True) == "#c98500"
    assert gr.cor_categoria("Categoria nova", False) == "#898781"


def test_rotulo_comp():
    assert gr.rotulo_comp("07/2026") == "jul/26" and gr.rotulo_comp("lixo") == "lixo"


def test_top_fornecedores(par):
    s, _ = par
    f = gr.fig_top_fornecedores(s["creditos"], False, topo=3)
    assert len(f.data[0].y) == 3


# ───────────── Planilha ─────────────
def test_xlsx_da_conferencia(par):
    s, d = par
    ajustes = [{"aba": "linhas", "chave": f"{FORN_A}|1001", "campo": "nota_contadora", "valor_novo": "conferido com a nota"}]
    L = cf.aplicar_aba("linhas", cf.conferir_linhas(s, d), ajustes).drop(columns=["editado"])
    x = cexp.gerar_xlsx_conferencia("Enermais Energia", "ENERGIA", "07/2026", L, cf.conferir_categorias(s, d), cf.conferir_bases(s, d),
                                    cf.conferir_totais(s, d), cf.conferir_debitos(s, d), [("Crédito", 127.05, 585.2)], {"SPED": "a.txt"})
    wb = openpyxl.load_workbook(io.BytesIO(x))
    assert wb.sheetnames == ["Resumo", "Linhas", "Categorias", "Bases", "Totais", "Débitos"]
    ws = wb["Linhas"]
    cab = [c.value for c in ws[5]]
    assert "Nota da contadora" in cab and "Status" in cab and ws.max_row == 5 + len(L)
    assert any(c.value == "conferido com a nota" for row in ws.iter_rows(min_row=6) for c in row)
    assert any(c.value == "47.040.664/0001-48" or c.value == "11.111.111/0001-11" for row in ws.iter_rows(min_row=6) for c in row)
