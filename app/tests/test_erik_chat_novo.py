# -*- coding: utf-8 -*-
"""Erik.AI v0.43.0 -- sem banco: prompt feminino, resposta em branco, propostas, visualizacao."""
import io
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import acoes_chat  # noqa: E402
import chat_egc  # noqa: E402
import chat_visual  # noqa: E402

EMPRESAS = ["ENERGIA", "SMG", "ENG", "RENOV", "CONST", "SOL"]


def _txt(t):
    return SimpleNamespace(type="text", text=t)


def _tool(nome, entrada, id="t1"):
    return SimpleNamespace(type="tool_use", name=nome, input=entrada, id=id)


def _client(*respostas):
    fila = list(respostas)
    chamadas = []

    def create(**kw):
        chamadas.append(kw)
        return fila.pop(0) if len(fila) > 1 else fila[0]

    return SimpleNamespace(messages=SimpleNamespace(create=create)), chamadas


def test_prompt_no_feminino_e_regras_novas():
    p = chat_egc.montar_system_prompt("a@b.com", EMPRESAS)
    assert "Voce e' a Erik.AI" in p and "Voce e' o Erik" not in p
    for trecho in ("FEMININO", "ANTI-ALUCINACAO", "NUNCA grava", "AGUARDANDO APROVACAO", "propor_acao"):
        assert trecho in p, trecho
    assert chat_egc.MAX_TOKENS >= 4096


def test_ferramentas_novas_estao_na_lista_e_no_dispatcher():
    nomes = {t["name"] for t in chat_egc.TOOLS}
    novas = {"consultar_evolucao_indicadores", "consultar_historico_importacoes", "consultar_nota_fiscal",
             "consultar_correcoes_manuais", "consultar_relatorios_gerados", "propor_acao"}
    assert novas <= nomes
    for t in chat_egc.TOOLS:
        assert t["input_schema"]["type"] == "object"
    r = chat_egc.executar_ferramenta(None, "ferramenta_que_nao_existe", {}, EMPRESAS)
    assert "erro" in r


def test_resposta_em_branco_faz_nova_tentativa_e_depois_usa_fallback():
    # 1a vazia, 2a (retry) com texto
    client, chamadas = _client(SimpleNamespace(stop_reason="end_turn", content=[_txt("")]),
                               SimpleNamespace(stop_reason="end_turn", content=[_txt("Resumo ok")]))
    r = chat_egc.responder(client, None, [{"role": "user", "content": "oi"}], "sys", EMPRESAS)
    assert r["texto"] == "Resumo ok" and len(chamadas) == 2
    assert chamadas[1]["messages"][-1]["role"] == "user" and "sem texto" in chamadas[1]["messages"][-1]["content"]
    # sempre vazia -> fallback explicito (nunca string vazia)
    client2, _ = _client(SimpleNamespace(stop_reason="end_turn", content=[]))
    r2 = chat_egc.responder(client2, None, [{"role": "user", "content": "oi"}], "sys", EMPRESAS)
    assert r2["texto"].strip() and "reformular" in r2["texto"]
    # vazia depois de ferramenta -> aponta pro painel
    client3, _ = _client(SimpleNamespace(stop_reason="tool_use", content=[_tool("consultar_periodos", {})]),
                         SimpleNamespace(stop_reason="end_turn", content=[]))
    with patch("consultas_chat.consultar_periodos", return_value={"periodos_por_empresa": {"SMG": ["2026-06"]}}):
        r3 = chat_egc.responder(client3, None, [{"role": "user", "content": "periodos?"}], "sys", EMPRESAS)
    assert "painel" in r3["texto"] and r3["ferramentas_usadas"][0]["nome"] == "consultar_periodos"


def test_max_tokens_avisa_corte():
    client, _ = _client(SimpleNamespace(stop_reason="max_tokens", content=[_txt("Resposta longa até aqui")]))
    r = chat_egc.responder(client, None, [{"role": "user", "content": "x"}], "sys", EMPRESAS)
    assert r["texto"].startswith("Resposta longa") and "cortada" in r["texto"]


def test_propor_acao_vira_proposta_e_o_modelo_so_ve_resumo_sem_executar():
    proposta = {"id": "abc12345", "tipo": "ARQUIVAR_PERIODO", "titulo": "Arquivar", "parametros": {"empresa": "SMG"},
                "resumo": [("Empresa", "SMG")], "editaveis": {}, "reversivel": True}
    client, chamadas = _client(
        SimpleNamespace(stop_reason="tool_use", content=[_tool("propor_acao", {"tipo": "ARQUIVAR_PERIODO", "empresa": "SMG", "periodo": "2026-06"})]),
        SimpleNamespace(stop_reason="end_turn", content=[_txt("Preparei; aguarda aprovação.")]))
    with patch.object(acoes_chat, "montar_proposta", return_value={"proposta": proposta}) as mp, \
         patch.object(acoes_chat, "executar_proposta") as ex:
        r = chat_egc.responder(client, None, [{"role": "user", "content": "arquive a SMG 06/2026"}], "sys", EMPRESAS)
    assert mp.called and not ex.called                       # a IA nunca executa
    assert r["propostas"] == [proposta] and r["ferramentas_usadas"] == []
    visto = chamadas[1]["messages"][-1]["content"][0]["content"]
    assert "NADA foi executado" in visto and "aguarda" in visto.lower() and "parametros" not in visto


def test_limite_de_propostas_por_resposta():
    p = lambda i: {"id": f"id{i}", "tipo": "ARQUIVAR_PERIODO", "titulo": "t", "parametros": {}, "resumo": [("a", "b")]}
    usos = [_tool("propor_acao", {"tipo": "ARQUIVAR_PERIODO"}, id=f"t{i}") for i in range(5)]
    client, _ = _client(SimpleNamespace(stop_reason="tool_use", content=usos), SimpleNamespace(stop_reason="end_turn", content=[_txt("ok")]))
    seq = iter(range(5))
    with patch.object(acoes_chat, "montar_proposta", side_effect=lambda *a, **k: {"proposta": p(next(seq))}):
        r = chat_egc.responder(client, None, [{"role": "user", "content": "x"}], "sys", EMPRESAS)
    assert len(r["propostas"]) == chat_egc.MAX_PROPOSTAS_POR_RESPOSTA


def test_acoes_tipo_fora_da_lista_e_valor_br():
    assert "erro" in acoes_chat.montar_proposta(None, "DELETE_TUDO", {}, EMPRESAS)
    assert acoes_chat._numero("R$ -1.234,50") == -1234.5 and acoes_chat._numero("1234.5") == 1234.5
    assert acoes_chat._numero("abc") is None and acoes_chat._numero(True) is None


# ───────────────── visualização ─────────────────

def test_visual_contas_visao_grupo_formatos_e_graficos():
    r = {"tipo": "BP", "periodo": "2026-06", "granularidade": "trimestral", "contas": [
        {"grupo": "AC", "conta": "CLIENTES", "VALOR CONSOLIDADO": 1500.5, "% ENERGIA": 0.66, "ENERMAIS ENERGIA": 1000.5},
        {"grupo": "PC", "conta": "FORN", "VALOR CONSOLIDADO": -300.0, "% ENERGIA": 0.5, "ENERMAIS ENERGIA": -150.0}]}
    prep = chat_visual.preparar("consultar_visao_grupo", r)
    t = chat_visual.tabela_exibicao(prep)
    assert t["VALOR CONSOLIDADO"].iloc[0] == "R$ 1.500,50" and "%" in t["% ENERGIA"].iloc[0]
    assert set(prep["graficos"]) == {"Total por grupo", "15 maiores contas"}
    for tipo in chat_visual.TIPOS_GRAFICO:
        fig = chat_visual.figura(prep["graficos"]["Total por grupo"], tipo, dark=False)
        assert len(fig.data) >= 1


def test_visual_excel_tem_numeros_reais_e_aba_fonte():
    import openpyxl
    r = {"evolucao": [{"periodo": "2026-03", "EBITDA": 10.0, "Margem EBITDA": 0.1}, {"periodo": "2026-06", "EBITDA": -5.0, "Margem EBITDA": -0.05}],
         "base_do_periodo": "Trimestral", "empresas_incluidas": ["ENERGIA"]}
    prep = chat_visual.preparar("consultar_evolucao_indicadores", r)
    wb = openpyxl.load_workbook(io.BytesIO(chat_visual.excel_bytes(prep, "u@x.com", "teste")))
    ws = wb["Resultado"]
    assert ws["B3"].value == -5.0 and isinstance(ws["B3"].value, (int, float))  # numero real, nao texto
    assert any("u@x.com" in str(c.value) for row in wb["Fonte"].iter_rows() for c in row)


def test_visual_pdf_gera_com_e_sem_grafico_e_erro_vira_none():
    prep = chat_visual.preparar("consultar_completude", {"completude_por_periodo": [
        {"Período": "2026-06", "Status": "✅ Completo (6/6)", "Empresas pendentes": "—"},
        {"Período": "2026-03", "Status": "❌ Faltam (4/6)", "Empresas pendentes": "SMG, SOL"}]})
    g = next(iter(prep["graficos"].values()))
    assert chat_visual.pdf_bytes(prep, "u", "c", g, "Barras")[:5] == b"%PDF-"
    assert chat_visual.pdf_bytes(prep, "u", "c")[:5] == b"%PDF-"
    assert chat_visual.preparar("consultar_bp_dre", {"erro": "x"}) is None
    assert chat_visual.preparar("consultar_bp_dre", {"contas": []}) is None


def test_visual_pendencias_nao_expoe_registro_id():
    r = {"pendencias": [{"registro_id": 7, "origem": "MANIFESTO", "numero_nota": "500", "status": "NAO_ENCONTRADA",
                         "fornecedor_nome": "A", "valor": 100.0}], "quantidade": 1}
    prep = chat_visual.preparar("consultar_notas_pendentes", r)
    assert "registro_id" not in prep["df"].columns
    assert chat_visual.tabela_exibicao(prep)["valor"].iloc[0] == "R$ 100,00"
