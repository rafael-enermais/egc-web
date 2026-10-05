# -*- coding: utf-8 -*-
import io
import sys
import datetime
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import nf_export  # noqa: E402


def _tabela():
    base = dict(cfop="5102", fornecedor_cnpj="07393522000140", confianca=None, observacao=None,
                pendencia_status=None, atualizado_em=datetime.datetime(2026, 9, 30, 10, 0, tzinfo=datetime.timezone.utc))
    return pd.DataFrame([
        dict(base, numero_nota="22299", data_emissao=datetime.date(2026, 7, 1), valor=4266.0,
             fornecedor_nome="SOLAR MOVEIS", status="LANCADA", confianca="NUMERO_CNPJ_VALOR",
             sienge_bill_id=32034, sienge_documento="NFE 22299", sienge_valor=4266.0, registro_id=1, origem="MANIFESTO"),
        dict(base, numero_nota="500", data_emissao=datetime.date(2026, 7, 5), valor=100.5,
             fornecedor_nome="OUTRO", status="NAO_ENCONTRADA", sienge_bill_id=None, sienge_documento=None,
             sienge_valor=None, pendencia_status="PENDENTE", registro_id=2, origem="MANIFESTO"),
        dict(base, numero_nota="[Sienge] nº 77", data_emissao=datetime.date(2026, 7, 9), valor=None,
             fornecedor_nome="ORFAO", status="SIENGE_SEM_MANIFESTO", sienge_bill_id=99, sienge_documento="NFE 77",
             sienge_valor=10.0, pendencia_status="PENDENTE", registro_id=3, origem="SIENGE_ORFAO"),
    ])


def test_preparar_tabela_renomeia_e_tira_internas():
    df = nf_export.preparar_tabela(_tabela())
    assert "registro_id" not in df.columns and "origem" not in df.columns
    assert "Título Sienge" in df.columns and "Documento no Sienge" in df.columns
    assert list(df.columns)[:3] == ["Nº da nota", "CFOP", "Título Sienge"]


def test_xlsx_completo_tem_todas_as_abas_e_todas_as_linhas():
    data = nf_export.gerar_xlsx_conferencia(_tabela(), "Enermais Energia", "07/2026", "manifesto.xlsx", "01/10/2026")
    xl = pd.ExcelFile(io.BytesIO(data))
    assert xl.sheet_names == ["Resumo", "Conferência completa", "Pendências", "Sienge sem nota"]
    comp = xl.parse("Conferência completa")
    assert len(comp) == 3  # lancada + pendente + orfao (nada some)
    assert comp.loc[comp["Nº da nota"].astype(str) == "22299", "Título Sienge"].iloc[0] == 32034
    assert comp.loc[comp["Nº da nota"].astype(str) == "22299", "Documento no Sienge"].iloc[0] == "NFE 22299"
    assert len(xl.parse("Pendências")) == 2
    assert len(xl.parse("Sienge sem nota")) == 1
    resumo = xl.parse("Resumo")
    assert resumo.loc[resumo["Item"] == "Notas no manifesto", "Valor"].iloc[0] == 2
    assert resumo.loc[resumo["Item"] == "Lançadas no Sienge", "Valor"].iloc[0] == 1


def test_xlsx_sem_orfaos_nem_pendencias_nao_quebra():
    t = _tabela().iloc[[0]]
    data = nf_export.gerar_xlsx_conferencia(t, "X", "07/2026")
    xl = pd.ExcelFile(io.BytesIO(data))
    assert len(xl.parse("Pendências")) == 0 and len(xl.parse("Sienge sem nota")) == 0


def test_aba_pendencias_traz_cfop_logo_depois_do_numero_01_10():
    import io
    import openpyxl
    t = _tabela()
    xlsx = nf_export.gerar_xlsx_conferencia(t, "Energia", "06/2026")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    cab = [c.value for c in wb["Pendências"][1]]
    assert cab[:2] == ["Nº da nota", "CFOP"], cab
    assert any(r[1] for r in wb["Pendências"].iter_rows(min_row=2, values_only=True)), "CFOP vazio nas pendencias"


# ───────────── v0.46.0: varios periodos numa conferencia so' ─────────────
def _tabela_2_periodos():
    import datetime as dt
    base = dict(cfop="5102", data_emissao=dt.date(2026, 7, 1), valor=10.0, fornecedor_nome="F", fornecedor_cnpj="1",
                confianca=None, sienge_bill_id=None, sienge_documento=None, sienge_valor=None, observacao=None,
                pendencia_status="PENDENTE", atualizado_em=dt.datetime(2026, 10, 1))
    L = [
        {**base, "periodo": "08/2026", "numero_nota": "3", "status": "NAO_ENCONTRADA", "registro_id": 3, "origem": "MANIFESTO"},
        {**base, "periodo": "07/2026", "numero_nota": "1", "status": "LANCADA", "registro_id": 1, "origem": "MANIFESTO"},
        {**base, "periodo": "07/2026", "numero_nota": "2", "status": "VALOR_DIVERGENTE", "registro_id": 2, "origem": "MANIFESTO"},
        {**base, "periodo": "08/2026", "numero_nota": "[Sienge] 9", "status": "SIENGE_SEM_MANIFESTO", "registro_id": 9, "origem": "SIENGE_ORFAO"},
    ]
    import pandas as pd
    return pd.DataFrame(L)


def test_rotulo_e_sufixo_de_periodos():
    import nf_export
    assert nf_export.rotulo_periodos(["07/2026"]) == "07/2026"
    assert nf_export.rotulo_periodos(["02/2026", "01/2026", "03/2026"]) == "01/2026 a 03/2026"
    assert nf_export.rotulo_periodos(["12/2025", "01/2026"]) == "12/2025 a 01/2026"       # vira o ano
    assert nf_export.rotulo_periodos(["01/2026", "03/2026"]) == "01/2026, 03/2026"        # com buraco: lista
    assert nf_export.rotulo_periodos([]) == ""
    assert nf_export.sufixo_arquivo(["07/2026"]) == "07_2026"
    assert nf_export.sufixo_arquivo(["01/2026", "02/2026"]) == "01_2026_a_02_2026"
    assert nf_export.sufixo_arquivo(["01/2026", "03/2026"]) == "01_2026_03_2026"


def test_resumo_por_periodo_soma_igual_ao_resumo_geral():
    import nf_export
    t = _tabela_2_periodos()
    rpp = nf_export.resumo_por_periodo(t)
    assert list(rpp["Período"]) == ["07/2026", "08/2026"]
    assert list(rpp["Notas no manifesto"]) == [2, 1] and list(rpp["Lançadas"]) == [1, 0]
    assert list(rpp["Sienge sem nota"]) == [0, 1] and list(rpp["Total de pendências"]) == [1, 2]
    geral = nf_export.resumo_conferencia(t)
    assert rpp["Total de pendências"].sum() == geral["pendencias"] and rpp["Notas no manifesto"].sum() == geral["total"]


def test_xlsx_de_varios_periodos_tem_resumo_por_periodo_e_coluna_periodo():
    import io
    import openpyxl
    import nf_export
    t = _tabela_2_periodos()
    wb = openpyxl.load_workbook(io.BytesIO(nf_export.gerar_xlsx_conferencia(t, "Energia", "07/2026 a 08/2026", "a.xlsx, b.xlsx", "x")))
    assert wb.sheetnames[:2] == ["Resumo", "Resumo por período"]
    ws = wb["Conferência completa"]
    cab = [c.value for c in ws[1]]
    assert cab[0] == "Período" and ws.max_row == 1 + len(t)
    assert wb["Pendências"].max_row == 1 + 3
    resumo = {r[0].value: r[1].value for r in wb["Resumo"].iter_rows(min_row=2)}
    assert resumo["Período de referência"] == "07/2026 a 08/2026" and resumo["Notas no manifesto"] == 3


def test_xlsx_de_um_periodo_continua_sem_a_aba_nova():
    import io
    import openpyxl
    import nf_export
    t = _tabela_2_periodos()
    t = t[t["periodo"] == "07/2026"]
    wb = openpyxl.load_workbook(io.BytesIO(nf_export.gerar_xlsx_conferencia(t, "Energia", "07/2026", None, "x")))
    assert "Resumo por período" not in wb.sheetnames
