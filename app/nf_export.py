# -*- coding: utf-8 -*-
"""
Exportacao da conferencia de Notas Fiscais x Sienge (tela 7_Notas_Fiscais).

01/10/2026 (Rafael: "eu baixo o resultado mas so' conseguimos baixar as
pendencias"): alem da planilha de pendencias, a contadora baixa a
conferencia COMPLETA -- todas as notas do manifesto (lancadas e
pendentes), o titulo e o documento correspondentes no Sienge, e os titulos
do Sienge sem nota. Funcoes puras (sem Streamlit/banco), testadas com
DataFrame sintetico em tests/test_nf_export.py.
"""
from __future__ import annotations

import io
from typing import Optional

import pandas as pd

import formatacao

# Ordem e rotulos das colunas na tela e nas planilhas. Colunas internas
# (registro_id, origem) nunca aparecem pra contadora.
ROTULOS = {
    "periodo": "Período",   # v0.46.0: so' existe quando a tabela junta varios periodos
    "numero_nota": "Nº da nota",
    "cfop": "CFOP",
    "sienge_bill_id": "Título Sienge",
    "sienge_documento": "Documento no Sienge",
    "status": "Status",
    "data_emissao": "Emissão",
    "valor": "Valor (manifesto)",
    "fornecedor_nome": "Fornecedor",
    "fornecedor_cnpj": "CNPJ do fornecedor",
    "confianca": "Critério do match",
    "sienge_valor": "Valor (Sienge)",
    "observacao": "Observação",
    "pendencia_status": "Acompanhamento",
    "atualizado_em": "Atualizado em",
}
ROTULOS_IGNORADAS = {
    "periodo": "Período",
    "numero_nota": "Nº da nota", "data_emissao": "Emissão", "valor": "Valor (manifesto)", "cfop": "CFOP",
    "fornecedor_nome": "Fornecedor", "fornecedor_cnpj": "CNPJ do fornecedor", "tipo_doc": "Tipo (Receita)",
    "natureza": "Natureza da operação", "motivo": "Motivo de ter ficado fora",
}
COLUNAS_MOEDA = ["Valor (manifesto)", "Valor (Sienge)"]
COLUNAS_DATA = ["Emissão"]
COLUNAS_INTERNAS = ["registro_id", "origem"]

# v0.44.3: rotulos amigaveis dos status (filtro, KPIs e Resumo do Excel).
# A coluna "Status" das tabelas continua com o codigo (e' o que a contadora ja' viu).
ROTULO_STATUS = {
    "LANCADA": "Lançadas",
    "LANCADA_OUTRA_EMPRESA": "Lançadas em outra empresa",
    "VALOR_DIVERGENTE": "Valor divergente",
    "NUMERO_DIVERGENTE": "Número divergente",
    "NAO_ENCONTRADA": "Não encontradas no Sienge",
    "SIENGE_SEM_MANIFESTO": "Sienge sem manifesto",
}
ORDEM_STATUS = list(ROTULO_STATUS)


def contar_por_status(tabela: pd.DataFrame) -> dict:
    """{status: n} na ordem de ORDEM_STATUS, incluindo zeros -- fonte unica
    pros KPIs da tela e pro Resumo do Excel (os dois sempre batem)."""
    contagem = tabela["status"].value_counts().to_dict() if len(tabela) and "status" in tabela.columns else {}
    resultado = {st: int(contagem.get(st, 0)) for st in ORDEM_STATUS}
    for st, n in contagem.items():  # status inesperado nunca some
        resultado.setdefault(st, int(n))
    return resultado


def resumo_conferencia(tabela: pd.DataFrame) -> dict:
    """Numeros do topo da conferencia. `pendencias` = tudo que NAO e'
    LANCADA (notas do manifesto + Sienge sem manifesto) = linhas da aba
    "Pendencias" do Excel e do botao "Baixar so as pendencias"."""
    manifesto = tabela[tabela["origem"] == "MANIFESTO"] if "origem" in tabela.columns else tabela
    orfaos = len(tabela) - len(manifesto)
    total = len(manifesto)
    lancadas = int((manifesto["status"] == "LANCADA").sum()) if total else 0
    return {
        "total": total, "lancadas": lancadas, "notas_pendentes": total - lancadas,
        "sienge_sem_manifesto": int(orfaos), "pendencias": int((tabela["status"] != "LANCADA").sum()) if len(tabela) else 0,
        "taxa": (lancadas / total) if total else None,
        "por_status": contar_por_status(tabela),
    }


def _chave_periodo(p) -> tuple:
    import re
    m = re.match(r"^\s*(\d{1,2})\s*/\s*(\d{4})\s*$", str(p or ""))
    return (int(m.group(2)), int(m.group(1))) if m else (9999, 99, str(p))


def periodos_ordenados(periodos) -> list:
    """Periodos unicos em ordem cronologica ('01/2026' < '02/2026' < ...)."""
    return sorted({str(p) for p in periodos if p}, key=_chave_periodo)


def rotulo_periodos(periodos) -> str:
    """'01/2026 a 08/2026' quando seguidos; senao '01/2026, 03/2026'; um so' -> '07/2026'."""
    ps = periodos_ordenados(periodos)
    if not ps:
        return ""
    if len(ps) == 1:
        return ps[0]
    ks = [_chave_periodo(p) for p in ps]
    if all(len(k) == 2 for k in ks):
        seguidos = all((b[0] * 12 + b[1]) - (a[0] * 12 + a[1]) == 1 for a, b in zip(ks, ks[1:]))
        if seguidos:
            return f"{ps[0]} a {ps[-1]}"
    return ", ".join(ps)


def sufixo_arquivo(periodos) -> str:
    """Parte do nome do arquivo: '07_2026' (um) | '01_2026_a_08_2026' (seguidos) | '01_2026_03_2026' (soltos)."""
    r = rotulo_periodos(periodos)
    if not r:
        return "periodo"
    return r.replace(" a ", "_a_").replace(", ", "_").replace("/", "_")


def resumo_por_periodo(tabela: pd.DataFrame) -> pd.DataFrame:
    """Uma linha por periodo (cronologico): notas, lancadas, pendentes, Sienge sem nota, total de pendencias, taxa.
    Vazio se a tabela nao tem a coluna `periodo`."""
    cols = ["Período", "Notas no manifesto", "Lançadas", "Notas com pendência", "Sienge sem nota",
            "Total de pendências", "Taxa de conciliação"]
    if tabela is None or tabela.empty or "periodo" not in tabela.columns:
        return pd.DataFrame(columns=cols)
    linhas = []
    for per in periodos_ordenados(tabela["periodo"]):
        r = resumo_conferencia(tabela[tabela["periodo"] == per])
        linhas.append([per, r["total"], r["lancadas"], r["notas_pendentes"], r["sienge_sem_manifesto"],
                       r["pendencias"], r["taxa"]])
    return pd.DataFrame(linhas, columns=cols)


def preparar_tabela(tabela: pd.DataFrame) -> pd.DataFrame:
    """Tira colunas internas, mantem a ordem de ROTULOS e renomeia pros
    rotulos de exibicao. Colunas desconhecidas sao descartadas."""
    cols = [c for c in ROTULOS if c in tabela.columns]
    return tabela[cols].rename(columns=ROTULOS).copy()


def preparar_ignoradas(ignoradas: pd.DataFrame) -> pd.DataFrame:
    """Notas fora da conferencia com rotulos de exibicao (colunas internas descartadas)."""
    cols = [c for c in ROTULOS_IGNORADAS if c in ignoradas.columns]
    return ignoradas[cols].rename(columns=ROTULOS_IGNORADAS).copy()


def _escrever_aba(writer, nome: str, df: pd.DataFrame) -> None:
    df = formatacao.remover_timezone_para_excel(df)
    if "Título Sienge" in df.columns:
        # inteiro sem separador de milhar; vazio quando nao ha' titulo
        df["Título Sienge"] = df["Título Sienge"].apply(lambda v: int(v) if pd.notna(v) else None).astype("object")
    for col in COLUNAS_DATA:
        if col in df.columns:
            df[col] = pd.to_datetime(df[col], errors="coerce").dt.date
    df.to_excel(writer, index=False, sheet_name=nome)
    ws = writer.sheets[nome]
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    cab = PatternFill("solid", fgColor="171C60")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = cab
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    for idx, col in enumerate(df.columns, start=1):
        letra = get_column_letter(idx)
        if col in COLUNAS_MOEDA:
            for c in ws[letra][1:]:
                c.number_format = '#,##0.00'
        if col in COLUNAS_DATA:
            for c in ws[letra][1:]:
                c.number_format = 'DD/MM/YYYY'
        tamanho = max([len(str(col))] + [len(str(v)) for v in df[col].head(300) if v is not None]) if len(df) else len(str(col))
        ws.column_dimensions[letra].width = min(max(tamanho + 2, 10), 48)


def gerar_xlsx_conferencia(
    tabela: pd.DataFrame, empresa_nome: str, periodo_referencia: str,
    arquivo_manifesto: Optional[str] = None, gerado_em: Optional[str] = None,
    ignoradas: Optional[pd.DataFrame] = None,
) -> bytes:
    """
    Planilha completa da conferencia. `tabela` = saida de
    nf_sienge.listar_conciliacao (+ listar_orfaos_sienge concatenado, com
    registro_id/origem). Abas: Resumo, Conferência completa (todas as
    notas do manifesto + titulos do Sienge sem nota), Pendências (tudo que
    nao e' LANCADA) e Sienge sem nota.
    v0.46.0: `tabela` pode juntar VARIOS periodos (coluna `periodo`): ganha a aba
    "Resumo por período" e a coluna Período em todas as abas; `periodo_referencia`
    e' o rotulo do conjunto (ex.: "01/2026 a 08/2026", ver rotulo_periodos).
    """
    orfaos = tabela[tabela["origem"] == "SIENGE_ORFAO"] if "origem" in tabela.columns else tabela.iloc[0:0]
    r = resumo_conferencia(tabela)

    resumo_linhas = [
        ("Empresa", empresa_nome),
        ("Período de referência", periodo_referencia),
        ("Arquivo do manifesto", arquivo_manifesto or "—"),
        ("Gerado em", gerado_em or "—"),
        ("Notas no manifesto", r["total"]),
        ("Lançadas no Sienge", r["lancadas"]),
        ("Taxa de conciliação", r["taxa"]),
        ("Notas do manifesto com pendência", r["notas_pendentes"]),
        ("Títulos no Sienge sem nota no manifesto", r["sienge_sem_manifesto"]),
        ("Total de pendências (aba Pendências)", r["pendencias"]),
        ("Detalhe por status", None),
    ] + [(f"  · {ROTULO_STATUS.get(st, st)}", n) for st, n in r["por_status"].items()]
    if ignoradas is not None and len(ignoradas):
        por_motivo = ignoradas["motivo"].value_counts().to_dict()
        resumo_linhas += [("Notas do arquivo que ficaram FORA da conferência (aba Ignoradas)", len(ignoradas))] + [
            (f"  · {m}", int(n)) for m, n in sorted(por_motivo.items())]
    resumo = pd.DataFrame(resumo_linhas, columns=["Item", "Valor"])
    por_periodo = resumo_por_periodo(tabela)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        resumo.to_excel(writer, index=False, sheet_name="Resumo")
        ws = writer.sheets["Resumo"]
        ws.column_dimensions["A"].width = 48
        ws.column_dimensions["B"].width = 36
        for row in ws.iter_rows(min_row=2):
            if row[0].value == "Taxa de conciliação" and row[1].value is not None:
                row[1].number_format = "0.0%"
        if len(por_periodo) > 1:
            por_periodo.to_excel(writer, index=False, sheet_name="Resumo por período")
            wsp = writer.sheets["Resumo por período"]
            for letra, larg in zip("ABCDEFG", (12, 20, 12, 22, 18, 20, 20)):
                wsp.column_dimensions[letra].width = larg
            for row in wsp.iter_rows(min_row=2):
                if row[6].value is not None:
                    row[6].number_format = "0.0%"
        _escrever_aba(writer, "Conferência completa", preparar_tabela(tabela))
        _escrever_aba(writer, "Pendências", preparar_tabela(tabela[tabela["status"] != "LANCADA"]))
        _escrever_aba(writer, "Sienge sem nota", preparar_tabela(orfaos))
        if ignoradas is not None and len(ignoradas):
            _escrever_aba(writer, "Ignoradas", preparar_ignoradas(ignoradas))
    return buf.getvalue()
