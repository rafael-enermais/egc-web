# -*- coding: utf-8 -*-
"""
EGC | Creditos PIS/COFINS -- planilha da conferencia (v0.50.0, 16/10/2026).

A contadora baixa a conferencia SPED x Detalhamento ja' com as edicoes dela aplicadas: abas Resumo, Linhas, Categorias, Bases, Totais e
Debitos. Mesmo padrao visual das outras planilhas (nf_export: logo, cabecalho da cor da marca, faixas, impressao). Funcoes puras.
"""
from __future__ import annotations

import io
from typing import Optional

import pandas as pd

import creditos_conferencia as cf
import formatacao
import nf_export

ROTULOS_LINHAS = {"nivel": "Nível", "cnpj": "CNPJ", "fornecedor": "Fornecedor", "numero": "Nº da nota", "categoria": "Categoria",
                  "pis_sped": "PIS (SPED)", "cofins_sped": "COFINS (SPED)", "pis_det": "PIS (Detalhamento)", "cofins_det": "COFINS (Detalhamento)",
                  "dif_pis": "Diferença PIS", "dif_cofins": "Diferença COFINS", "status": "Status", "observacao": "Observação",
                  "nota_contadora": "Nota da contadora", "conferido": "Conferido"}
ROTULOS_CATEGORIAS = {"categoria": "Categoria", "pis_sped": "PIS (SPED)", "cofins_sped": "COFINS (SPED)", "pis_det": "PIS (Detalhamento)",
                      "cofins_det": "COFINS (Detalhamento)", "linhas_det": "Linhas no Detalhamento", "dif_pis": "Diferença PIS",
                      "dif_cofins": "Diferença COFINS", "status": "Status"}
ROTULOS_BASES = {"grupo": "Grupo", "base_sped": "Base (SPED)", "base_det": "Base (Detalhamento)", "dif": "Diferença", "status": "Status",
                 "pis_pela_base": "PIS pela base", "pis_linhas": "PIS pelas linhas", "dif_interna_pis": "Diferença interna PIS",
                 "cofins_pela_base": "COFINS pela base", "cofins_linhas": "COFINS pelas linhas", "dif_interna_cofins": "Diferença interna COFINS"}
ROTULOS_TOTAIS = {"tributo": "Tributo", "credito_bloco_m": "Crédito no Bloco M", "soma_detalhamento": "Soma do Detalhamento",
                  "depreciacao_imobilizado_sped": "Depreciação/Imobilizado (SPED)", "inconsistencia_sped": "Inconsistência dentro do SPED",
                  "residuo": "Resíduo", "status": "Status", "observacao": "Observação"}
ROTULOS_DEBITOS = {"item": "Item da apuração", "pis_sped": "PIS (SPED)", "pis_det": "PIS (Detalhamento)", "dif_pis": "Diferença PIS",
                   "cofins_sped": "COFINS (SPED)", "cofins_det": "COFINS (Detalhamento)", "dif_cofins": "Diferença COFINS", "status": "Status",
                   "nota_contadora": "Nota da contadora"}
_MOEDA = {"PIS (SPED)", "COFINS (SPED)", "PIS (Detalhamento)", "COFINS (Detalhamento)", "Diferença PIS", "Diferença COFINS", "Base (SPED)",
          "Base (Detalhamento)", "Diferença", "PIS pela base", "PIS pelas linhas", "Diferença interna PIS", "COFINS pela base",
          "COFINS pelas linhas", "Diferença interna COFINS", "Crédito no Bloco M", "Soma do Detalhamento", "Depreciação/Imobilizado (SPED)",
          "Inconsistência dentro do SPED", "Resíduo"}
_COR_STATUS = {cf.S_OK: "E1F5EE", cf.S_ESPERADO: "E6F1FB", cf.S_NUM: "FAEEDA", cf.S_DIF: "FCEBEB", cf.S_SO_SPED: "FCEBEB", cf.S_SO_DET: "FCEBEB",
               cf.S_MANUAL: "EEEDFE"}
_LARGURA = {"Fornecedor": 36, "Observação": 46, "Nota da contadora": 34, "Item da apuração": 56, "Grupo": 44, "Categoria": 26, "CNPJ": 20,
            "Status": 22, "Nível": 16}


def _preparar(df: pd.DataFrame, rotulos: dict) -> pd.DataFrame:
    d = df[[c for c in rotulos if c in df.columns]].copy()
    if "cnpj" in d.columns:
        d["cnpj"] = d["cnpj"].map(nf_export._cnpj_formatado)
    if "conferido" in d.columns:
        d["conferido"] = d["conferido"].map(lambda v: "Sim" if bool(v) else "")
    return d.rename(columns=rotulos)


def _escrever(writer, nome: str, df: pd.DataFrame, titulo: str, subtitulo: str, empresa_codigo: Optional[str]) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    pal = nf_export.paleta_da_empresa(empresa_codigo)
    L = nf_export.LINHA_CABECALHO
    df.to_excel(writer, index=False, sheet_name=nome, startrow=L - 1)
    ws = writer.sheets[nome]
    n_col = max(len(df.columns), 1)
    nf_export._topo(ws, titulo, subtitulo, empresa_codigo, n_col, col_titulo=min(3, n_col))
    for cell in ws[L]:
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=pal["primaria"])
        cell.border = Border(bottom=Side(style="medium", color=pal["destaque"]))
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[L].height = 30
    borda = Border(bottom=Side(style="thin", color=nf_export.CINZA_BORDA))
    zebra = PatternFill("solid", fgColor=nf_export.CINZA_ZEBRA)
    cols = {c: i for i, c in enumerate(df.columns, start=1)}
    for k in range(len(df)):
        r = L + 1 + k
        for nome_col, i in cols.items():
            c = ws.cell(row=r, column=i)
            c.font = Font(name="Calibri", size=10)
            c.border = borda
            if k % 2:
                c.fill = zebra
            if nome_col in _MOEDA:
                c.number_format = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
                c.alignment = Alignment(horizontal="right", vertical="top")
            elif nome_col in ("Observação", "Nota da contadora", "Item da apuração"):
                c.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                c.alignment = Alignment(horizontal="left", vertical="top")
        if "Status" in cols:
            cor = _COR_STATUS.get(df.iat[k, cols["Status"] - 1])
            if cor:
                c = ws.cell(row=r, column=cols["Status"])
                c.fill = PatternFill("solid", fgColor=cor)
                c.font = Font(name="Calibri", size=10, bold=True, color=pal["primaria"])
    for nome_col, i in cols.items():
        ws.column_dimensions[get_column_letter(i)].width = _LARGURA.get(nome_col, 16)
    ws.freeze_panes = ws.cell(row=L + 1, column=1)
    ws.auto_filter.ref = f"A{L}:{get_column_letter(n_col)}{L + max(len(df), 1)}"
    nf_export._formatar_impressao(ws, L, n_col, empresa_codigo)


def _aba_resumo(writer, empresa_nome: str, empresa_codigo: Optional[str], competencia: str, kpis: list, status: pd.DataFrame,
                arquivos: dict) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    pal = nf_export.paleta_da_empresa(empresa_codigo)
    ws = writer.book.create_sheet("Resumo", 0)
    nf_export._topo(ws, "Créditos PIS/COFINS — conferência SPED × Detalhamento", f"{empresa_nome} · competência {competencia} · gerado em {formatacao.agora_br()}",
                    empresa_codigo, 6)
    r = 6
    for rot, val in [("Empresa", empresa_nome), ("Competência", competencia)] + [(k, v) for k, v in arquivos.items()]:
        ws.cell(row=r, column=1, value=rot).font = Font(name="Calibri", size=10, bold=True, color=pal["primaria"])
        ws.cell(row=r, column=2, value=val).font = Font(name="Calibri", size=10)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Números do mês").font = Font(name="Calibri", size=12, bold=True, color=pal["primaria"])
    r += 1
    for rot, pis, cof in kpis:
        ws.cell(row=r, column=1, value=rot).font = Font(name="Calibri", size=10)
        for j, v in ((2, pis), (3, cof)):
            c = ws.cell(row=r, column=j, value=v)
            c.number_format = '"R$" #,##0.00'
            c.font = Font(name="Calibri", size=10)
            c.alignment = Alignment(horizontal="right")
        r += 1
    ws.cell(row=r - len(kpis) - 1, column=2, value="PIS").font = Font(name="Calibri", size=10, bold=True)
    ws.cell(row=r - len(kpis) - 1, column=3, value="COFINS").font = Font(name="Calibri", size=10, bold=True)
    r += 1
    ws.cell(row=r, column=1, value="Conferência das linhas").font = Font(name="Calibri", size=12, bold=True, color=pal["primaria"])
    r += 1
    for s in status.itertuples(index=False):
        c = ws.cell(row=r, column=1, value=s.status)
        c.font = Font(name="Calibri", size=10, bold=True, color=pal["primaria"])
        cor = _COR_STATUS.get(s.status)
        if cor:
            c.fill = PatternFill("solid", fgColor=cor)
        ws.cell(row=r, column=2, value=int(s.linhas)).font = Font(name="Calibri", size=10)
        r += 1
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 20
    nf_export._formatar_impressao(ws, 5, 3, empresa_codigo)


def gerar_xlsx_conferencia(empresa_nome: str, empresa_codigo: Optional[str], competencia: str, linhas: pd.DataFrame, categorias: pd.DataFrame,
                           bases: pd.DataFrame, totais: pd.DataFrame, debitos: pd.DataFrame, kpis: list,
                           arquivos: Optional[dict] = None) -> bytes:
    """Devolve o .xlsx. `kpis` = [(rotulo, pis, cofins)]; `arquivos` = {'SPED': nome, 'Detalhamento': nome} (so' informativo)."""
    buf = io.BytesIO()
    sub = f"{empresa_nome} · {competencia}"
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        _escrever(w, "Linhas", _preparar(linhas, ROTULOS_LINHAS), "Conferência linha a linha", sub, empresa_codigo)
        _escrever(w, "Categorias", _preparar(categorias, ROTULOS_CATEGORIAS), "Crédito por categoria", sub, empresa_codigo)
        _escrever(w, "Bases", _preparar(bases, ROTULOS_BASES), "Base de cálculo por natureza", sub, empresa_codigo)
        _escrever(w, "Totais", _preparar(totais, ROTULOS_TOTAIS), "Crédito do Bloco M × Detalhamento", sub, empresa_codigo)
        _escrever(w, "Débitos", _preparar(debitos, ROTULOS_DEBITOS), "Apuração (débitos e créditos utilizados)", sub, empresa_codigo)
        _aba_resumo(w, empresa_nome, empresa_codigo, competencia, kpis, cf.resumo_linhas(linhas), arquivos or {})
    return buf.getvalue()
